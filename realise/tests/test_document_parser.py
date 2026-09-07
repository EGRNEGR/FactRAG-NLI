"""Real document-adapter and source preservation regression tests."""

import os
from pathlib import Path

import pytest

from document_parser import ChunkingConfig, DocumentParseError, DocumentParser


def test_long_paragraph_is_bounded() -> None:
    parser = DocumentParser(ChunkingConfig(target_tokens=60, max_tokens=90, overlap_tokens=10))
    parsed = parser.parse_text("# Раздел\n" + "Давление: 10,5 МПа; " * 100)
    assert len(parsed.chunks) > 2
    assert all(chunk.token_count <= 90 for chunk in parsed.chunks)
    assert any("10,5 МПа" in chunk.text for chunk in parsed.chunks)
    assert not any(chunk.metadata["oversized"] for chunk in parsed.chunks)


def test_numbered_hierarchy_and_footnote() -> None:
    parsed = DocumentParser().parse_text(
        "4 Требования\n4.2 Давление\n4.2.1 Проверка\nДопуск 5 [^1].\n[^1]: При 20 °C."
    )
    target = next(e for e in parsed.elements if e.text.startswith("Допуск"))
    assert target.section_path == ("4 Требования", "4.2 Давление", "4.2.1 Проверка")
    assert target.metadata["footnote_context"] == ("[^1]: При 20 °C.",)


def test_table_rows_remain_atomic_and_repeat_header() -> None:
    rows = [f"| item{i} | value{i} |" for i in range(40)]
    parsed = DocumentParser(
        ChunkingConfig(target_tokens=40, max_tokens=70, overlap_tokens=5)
    ).parse_text("Таблица 2 — Параметры\n| Name | Value |\n| --- | --- |\n" + "\n".join(rows))
    chunks = [c for c in parsed.chunks if "table" in c.metadata["kinds"]]
    assert len(chunks) > 1
    for row in rows:
        assert sum(row in c.text for c in chunks) == 1
    assert all("| Name | Value |" in c.text for c in chunks)
    assert all(c.metadata["table_context"][0]["table_number"] == "2" for c in chunks)


def test_continuation_requires_same_number() -> None:
    parsed = DocumentParser().parse_text(
        "Таблица 1\n| A | B |\n| 1 | 2 |\n\nПродолжение таблицы 1\n| A | B |\n| 3 | 4 |\n\nПродолжение таблицы 2\n| A | B |\n| 5 | 6 |"
    )
    tables = [e for e in parsed.elements if e.kind == "table"]
    assert tables[1].metadata["continuation_of"] == tables[0].element_id
    assert tables[2].metadata["continuation_unresolved"]


def test_content_id_ignores_mtime_and_path(tmp_path: Path) -> None:
    first, second = tmp_path / "first.txt", tmp_path / "second.txt"
    first.write_text("# A\nText", encoding="utf-8")
    second.write_bytes(first.read_bytes())
    os.utime(first, (1, 1))
    one, two = DocumentParser().parse(first), DocumentParser().parse(second)
    assert one.document_id == two.document_id
    assert [c.chunk_id for c in one.chunks] == [c.chunk_id for c in two.chunks]


def test_docx_order_and_merged_cells(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Requirements", level=1)
    document.add_paragraph("Before")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Shared header"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    table.cell(2, 0).text = "10"
    table.cell(2, 1).text = "20"
    document.add_paragraph("After")
    path = tmp_path / "source.docx"
    document.save(str(path))
    parsed = DocumentParser().parse(path)
    assert [e.kind for e in parsed.elements] == ["heading", "paragraph", "table", "paragraph"]
    assert "Shared header | Shared header" in parsed.elements[2].text
    assert parsed.elements[2].page is None


def test_pdf_reading_order_and_page_limit(tmp_path: Path) -> None:
    import pymupdf

    path = tmp_path / "source.pdf"
    with pymupdf.open() as document:  # type: ignore[no-untyped-call]  # PyMuPDF factory lacks annotations.
        page = document.new_page()
        page.insert_text((50, 100), "4 Requirements")
        page.insert_text((50, 140), "Before table")
        for x in (50, 150, 250):
            page.draw_line((x, 180), (x, 260))
        for y in (180, 220, 260):
            page.draw_line((50, y), (250, y))
        for x, y, text in (
            (60, 200, "Name"),
            (160, 200, "Value"),
            (60, 240, "Pressure"),
            (160, 240, "10"),
        ):
            page.insert_text((x, y), text)
        page.insert_text((50, 300), "After table")
        document.save(path)
    parsed = DocumentParser().parse(path)
    assert [e.kind for e in parsed.elements] == ["heading", "paragraph", "table", "paragraph"]
    assert parsed.elements[2].section_path == ("4 Requirements",)
    assert parsed.elements[2].page == 1


def test_empty_pdf_fails_explicitly(tmp_path: Path) -> None:
    import pymupdf

    path = tmp_path / "blank.pdf"
    with pymupdf.open() as document:  # type: ignore[no-untyped-call]  # PyMuPDF factory lacks annotations.
        document.new_page()
        document.save(path)
    with pytest.raises(DocumentParseError, match="OCR required"):
        DocumentParser().parse(path)


def test_docx_multilevel_header(tmp_path: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement

    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Pressure"
    table.cell(1, 0).text = "Minimum"
    table.cell(1, 1).text = "Maximum"
    table.cell(2, 0).text = "5"
    table.cell(2, 1).text = "10"
    for row in table.rows[:2]:
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    path = tmp_path / "headers.docx"
    document.save(str(path))
    parsed = DocumentParser().parse(path)
    assert "| Pressure / Minimum | Pressure / Maximum |" in parsed.elements[0].text
    assert "| 5 | 10 |" in parsed.elements[0].text


def test_docx_native_footnote(tmp_path: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from zipfile import ZipFile

    document = Document()
    paragraph = document.add_paragraph("Limit 10")
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), "1")
    paragraph.add_run()._r.append(reference)
    path = tmp_path / "note.docx"
    document.save(str(path))
    with ZipFile(path, "a") as archive:
        archive.writestr(
            "word/footnotes.xml",
            '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Only at 20 C</w:t></w:r></w:p></w:footnote></w:footnotes>',
        )
    parsed = DocumentParser().parse(path)
    assert parsed.elements[0].metadata["footnote_context"] == ("[^1]: Only at 20 C",)


def test_default_parser_fails_without_production_weights(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("RAG_MODE", "production")
    monkeypatch.setenv("RAG_ALLOW_FALLBACK", "false")
    monkeypatch.setenv("RAG_API_TOKENS", '["01234567890123456789012345678901"]')
    monkeypatch.setenv("RAG_EMBEDDING_MODEL_PATH", "missing-model-for-test")
    with pytest.raises(ValueError, match="directory does not exist"):
        DocumentParser()
