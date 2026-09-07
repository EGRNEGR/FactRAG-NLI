"""Real Qdrant/BM25 persistence with explicit development embedding adapters."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Sequence

import pytest

from document_parser import DocumentChunk, DocumentParseError, RegexTokenCounter
from document_processor import DocumentProcessor, index_status
from hybrid_retriever import HashingEncoder, HybridRetriever, LexicalReranker, RetrievalError
from scripts.manage_index import main
from settings import RAGSettings


def config(tmp_path: Path) -> RAGSettings:
    return RAGSettings(
        mode="development",
        allow_fallback=True,
        qdrant_path=tmp_path / "index",
        rerank_score_threshold=0.1,
    )


def long_text() -> str:
    return "# Network\n\n" + "\n\n".join(
        " ".join(
            f"Sentence {i * 12 + j} describes reliable network packet delivery." for j in range(12)
        )
        for i in range(40)
    )


def search(settings: RAGSettings) -> HybridRetriever:
    return HybridRetriever(HashingEncoder(64), LexicalReranker(), settings=settings)


def test_whole_sentences_and_real_overlap(tmp_path: Path) -> None:
    processor = DocumentProcessor(config(tmp_path), tokenizer=RegexTokenCounter())
    parsed = processor.parser.parse_text(long_text(), document_id="doc")
    assert len(parsed.chunks) > 3
    texts = [c.text.split("\n", 2)[2] for c in parsed.chunks]
    assert all(c.token_count <= 800 for c in parsed.chunks)
    for previous, current in zip(texts, texts[1:]):
        a, b = previous.splitlines(), current.splitlines()
        overlap = max((n for n in range(1, min(len(a), len(b)) + 1) if a[-n:] == b[:n]), default=0)
        assert overlap > 0
        assert 50 <= RegexTokenCounter().count("\n".join(b[:overlap])) <= 75
        assert all(line.endswith(".") for line in current.splitlines())
    for i in range(480):
        assert any(f"Sentence {i} describes reliable network packet delivery." in t for t in texts)


def test_dedup_before_parser_and_after_restart(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = config(tmp_path)
    source = tmp_path / "one.md"
    source.write_text("# Pressure\n\nPressure is 10 MPa.", encoding="utf-8")
    renamed = tmp_path / "renamed.txt"
    renamed.write_bytes(source.read_bytes())
    with search(settings) as index:
        first = DocumentProcessor(settings).add(source, index)
        assert first.status == "added"
        assert first.sha256 == hashlib.sha256(source.read_bytes()).hexdigest()
    with search(settings) as index:
        processor = DocumentProcessor(settings)
        monkeypatch.setattr(processor, "_parse", lambda *args: pytest.fail("Duplicate was parsed"))
        result = processor.add(renamed, index)
        assert result.status == "duplicate"
        assert index.size == first.chunks
    assert index_status(settings).documents == 1


class RecordingIndex:
    def __init__(self) -> None:
        self.chunks: tuple[DocumentChunk, ...] = ()
        self.batch_size = 0

    def document_chunks(self, document_id: str) -> tuple[DocumentChunk, ...]:
        return ()

    def add_incremental(self, chunks: Sequence[DocumentChunk], *, batch_size: int) -> int:
        self.chunks, self.batch_size = tuple(chunks), batch_size
        return len(chunks)

    def delete_document(self, document_id: str) -> int:
        pytest.fail("New file must not delete anything")


def test_mock_incremental_manifest_and_source(tmp_path: Path) -> None:
    source = tmp_path / "new.md"
    source.write_text(long_text(), encoding="utf-8")
    index = RecordingIndex()
    result = DocumentProcessor(config(tmp_path)).add(source, index)
    assert index.batch_size == 2
    assert len(index.chunks) == result.chunks
    for i, chunk in enumerate(index.chunks):
        assert chunk.source == str(source.resolve())
        assert chunk.metadata["ingestion_chunk_count"] == result.chunks
        assert chunk.metadata["ingestion_ordinal"] == i
        assert chunk.document_id == chunk.metadata["file_sha256"] == result.sha256


def test_docx_preserves_paragraph_order_and_indexes(tmp_path: Path) -> None:
    """Мы извлекаем настоящий Word-файл и передаём его абзацы в общий индекс."""
    from docx import Document
    from web_bridge import safe_filename

    source = tmp_path / "manual.DOCX"
    document = Document()
    document.add_heading("Наш протокол", level=1)
    paragraphs = ["Первый абзац: давление равно 10 МПа.", "Второй абзац: температура равна 25 °C."]
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(source))

    processor = DocumentProcessor(config(tmp_path), tokenizer=RegexTokenCounter())
    parsed = processor.parse_docx(source)
    assert [e.text for e in parsed.elements if e.kind == "paragraph"] == paragraphs
    index = RecordingIndex()
    result = processor.add(source, index)
    assert result.status == "added"
    assert result.sha256 == hashlib.sha256(source.read_bytes()).hexdigest()
    assert index.chunks
    text = "\n".join(c.text for c in index.chunks)
    assert text.index(paragraphs[0]) < text.index(paragraphs[1])
    assert all(c.source == str(source.resolve()) for c in index.chunks)
    assert any("Наш протокол" in c.section_path for c in index.chunks)
    assert safe_filename(source.name) == "manual.docx"


def test_append_preserves_existing_retrieval(tmp_path: Path) -> None:
    settings = config(tmp_path)
    a, b = tmp_path / "a.txt", tmp_path / "b.txt"
    a.write_text("Pressure is 10 MPa.", encoding="utf-8")
    b.write_text("Temperature is 300 Kelvin.", encoding="utf-8")
    with search(settings) as index:
        processor = DocumentProcessor(settings)
        first = processor.add(a, index)
        ids = [h.chunk.chunk_id for h in index.search("Pressure")]
        second = processor.add(b, index)
        assert index.size == first.chunks + second.chunks
        assert ids[0] == index.search("Pressure")[0].chunk.chunk_id
        assert index.search("Temperature")[0].chunk.document_id == second.sha256
    status = index_status(settings)
    assert status.documents == 2 and status.incomplete_documents == 0


def test_partial_write_hidden_and_recovered(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = config(tmp_path)
    source = tmp_path / "large.md"
    source.write_text(long_text(), encoding="utf-8")
    processor = DocumentProcessor(settings)
    with search(settings) as index:
        encode = index.encoder.encode
        calls = 0

        def fail_second(texts: Sequence[str]) -> list[list[float]]:
            nonlocal calls
            calls += 1
            assert len(texts) <= 2
            if calls == 2:
                raise RuntimeError("Simulated GPU failure")
            return encode(texts)

        monkeypatch.setattr(index.encoder, "encode", fail_second)
        with pytest.raises(RetrievalError, match="reopen"):
            processor.add(source, index)
    status = index_status(settings)
    assert status.documents == 0 and status.incomplete_documents == 1 and status.stored_chunks == 2
    with search(settings) as index:
        assert index.search("network packet") == []
        recovered = processor.add(source, index)
        assert recovered.status == "added"
        assert index.search("network packet")
    assert index_status(settings).incomplete_documents == 0


def test_legacy_hash_is_already_indexed(tmp_path: Path) -> None:
    settings = config(tmp_path)
    path = tmp_path / "legacy.txt"
    path.write_text("Pressure is 10 MPa.", encoding="utf-8")
    from document_parser import DocumentParser

    parsed = DocumentParser(tokenizer=RegexTokenCounter()).parse(path)
    with search(settings) as index:
        index.add(parsed.chunks)
        assert DocumentProcessor(settings).add(path, index).status == "duplicate"


@pytest.mark.parametrize("text", ["", "\x00\u200b", "word " * 1000 + "."])
def test_invalid_text_does_not_write(tmp_path: Path, text: str) -> None:
    path = tmp_path / "invalid.txt"
    path.write_text(text, encoding="utf-8")
    index = RecordingIndex()
    with pytest.raises(DocumentParseError):
        DocumentProcessor(config(tmp_path)).add(path, index)
    assert not index.chunks


def test_pdf_text_and_page_metadata(tmp_path: Path) -> None:
    import pymupdf

    path = tmp_path / "sample.pdf"
    with pymupdf.open() as pdf:  # type: ignore[no-untyped-call]
        page = pdf.new_page()
        page.insert_text((72, 200), "Network pressure is 10 MPa.")
        pdf.save(path)
    index = RecordingIndex()
    DocumentProcessor(config(tmp_path)).add(path, index)
    assert "Network pressure is 10 MPa." in index.chunks[0].text
    assert index.chunks[0].page_start == index.chunks[0].page_end == 1


def test_status_cli_does_not_load_models(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    import scripts.manage_index as cli

    monkeypatch.setattr(cli, "RAGSettings", lambda: config(tmp_path))
    monkeypatch.setattr(cli, "HybridRetriever", lambda *a, **k: pytest.fail("Model loading"))
    assert main(["--status"]) == 0
    assert '"documents": 0' in capsys.readouterr().out
    assert not (tmp_path / "index").exists()


def test_cli_scan_continues_after_bad_file_and_add_skips_duplicate(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    import scripts.manage_index as cli

    settings = config(tmp_path)
    docs = tmp_path / "documents"
    docs.mkdir()
    (docs / "a.txt").write_text("Pressure is 10 MPa.", encoding="utf-8")
    (docs / "b.md").write_text("Temperature is 300 Kelvin.", encoding="utf-8")
    (docs / "bad.txt").write_bytes(b"\xff\xfe\xff")
    monkeypatch.setattr(cli, "RAGSettings", lambda: settings)
    monkeypatch.setattr(cli, "HybridRetriever", lambda *, settings: search(settings))
    assert main(["--scan", str(docs)]) == 1
    assert '"status": "error"' in capsys.readouterr().out
    assert index_status(settings).documents == 2
    assert main(["--add", str(docs / "a.txt")]) == 0
    assert '"status": "duplicate"' in capsys.readouterr().out
