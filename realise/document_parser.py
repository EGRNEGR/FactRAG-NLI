"""Hierarchical parsing and semantic chunking for technical documents.

The parser keeps the source structure in a normalized intermediate representation.
Heavy document libraries are optional and imported only when their adapter is used,
which makes the module suitable for an offline on-premise deployment and unit tests.
"""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import TYPE_CHECKING, Any, Iterable, Iterator, Mapping, Sequence, Protocol

if TYPE_CHECKING:
    from settings import RAGSettings


class DocumentParseError(RuntimeError):
    """Raised when a document cannot be parsed by the selected backend."""


@dataclass(slots=True, frozen=True)
class DocumentElement:
    """A normalized source element preserving order and document hierarchy."""

    element_id: str
    kind: str  # heading, paragraph, list, table, footnote, caption
    text: str
    level: int = 0
    number: str | None = None
    section_path: tuple[str, ...] = ()
    page: int | None = None
    source: str | None = None
    metadata: Mapping[str, Any] = field(default_factory=dict)


@dataclass(slots=True, frozen=True)
class DocumentChunk:
    """Retrieval unit with stable citation metadata."""

    chunk_id: str
    text: str
    document_id: str
    source: str
    section_path: tuple[str, ...]
    page_start: int | None
    page_end: int | None
    element_ids: tuple[str, ...]
    token_count: int
    metadata: Mapping[str, Any] = field(default_factory=dict)

    @property
    def citation(self) -> str:
        """Return a human-readable citation suitable for an answer."""
        path = " > ".join(self.section_path) or self.document_id
        pages = f", стр. {self.page_start}" if self.page_start else ""
        if self.page_end and self.page_end != self.page_start:
            pages = f", стр. {self.page_start}-{self.page_end}"
        return f"{path}{pages}"


ParsedChunk = DocumentChunk


@dataclass(slots=True, frozen=True)
class ParsedDocument:
    """Parsed document and its semantically chunked representation."""

    document_id: str
    source: str
    elements: tuple[DocumentElement, ...]
    chunks: tuple[DocumentChunk, ...]
    metadata: Mapping[str, Any] = field(default_factory=dict)


@dataclass(slots=True, frozen=True)
class ChunkingConfig:
    """Configuration for the adaptive section-aware chunker."""

    target_tokens: int = 420
    max_tokens: int = 600
    overlap_tokens: int = 70
    min_tokens: int = 40
    keep_tables_intact: bool = True

    def __post_init__(self) -> None:
        if not 0 <= self.overlap_tokens < self.max_tokens:
            raise ValueError("Require 0 <= overlap < max_tokens")
        if not 0 < self.target_tokens <= self.max_tokens or self.min_tokens < 0:
            raise ValueError("Require 0 < target <= max_tokens and nonnegative minimum")


@dataclass(slots=True, frozen=True)
class ChunkMetadata:
    """Explicit source context attached to a retrieval unit."""

    document_id: str
    chunk_id: str
    section_path: tuple[str, ...]
    page_start: int | None
    page_end: int | None
    table_context: tuple[Mapping[str, Any], ...]
    footnote_context: tuple[str, ...]
    token_count: int


class TokenCounter(Protocol):
    """Count tokens without altering source text."""

    def count(self, text: str) -> int:
        """Return token count, excluding model special tokens."""
        ...


class LocalTokenizer:
    """Use an installed model tokenizer exclusively from local artifacts."""

    def __init__(self, path: Path) -> None:
        from transformers import AutoTokenizer

        self._tokenizer = AutoTokenizer.from_pretrained(
            str(path.resolve(strict=True)),
            local_files_only=True,
            trust_remote_code=False,
        )

    def count(self, text: str) -> int:
        return len(self._tokenizer.encode(text, add_special_tokens=False))


class RegexTokenCounter:
    """Explicit lexical counter for development and parser-only tests."""

    def count(self, text: str) -> int:
        return len(_tokens(text))


_NUMBERED_HEADING = re.compile(r"^\s*(\d+(?:\.\d+)*\.?|[A-ZА-Я]\.)\s+(.+?)\s*$")
_MARKDOWN_HEADING = re.compile(r"^\s*(#{1,6})\s+(.+?)\s*$")
_FOOTNOTE = re.compile(
    r"^\s*(?:\[\^?\d+\]|\d+\)|Примечание(?:\s*[-—:]|$)|NOTE(?:\s*[-—:]|$))", re.I
)
_TOKEN = re.compile(r"\w+|[^\w\s]", re.UNICODE)
_CAPTION = re.compile(r"^(?:(Продолжение|Окончание)\s+)?[Тт]аблиц[аы]\s+([\w.\-]+)(.*)$")
_MARKER = re.compile(r"^\s*(?:\[\^?(\d+)\][:]?|(\d+)\)|(\*+))\s*(.*)$")


def _tokens(text: str) -> list[str]:
    return _TOKEN.findall(text)


def _stable_id(*parts: str) -> str:
    return hashlib.sha1("\x1f".join(parts).encode("utf-8")).hexdigest()[:16]


def _looks_like_table(line: str, next_line: str | None) -> bool:
    return "|" in line and (line.count("|") >= 2 or (next_line is not None and "|" in next_line))


def _table_to_markdown(lines: Sequence[str]) -> str:
    """Normalize pipe/tab-separated rows into a complete Markdown table."""
    rows: list[list[str]] = []
    for raw in lines:
        raw = raw.strip()
        if not raw:
            continue
        if "|" in raw:
            cells = [cell.strip() for cell in raw.strip("|").split("|")]
        else:
            cells = [cell.strip() for cell in raw.split("\t")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines_out = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines_out.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines_out)


class DocumentParser:
    """Deterministic local adapters with source-preserving, bounded chunking.

    Without an injected tokenizer, environment settings enforce production checks.
    A lexical counter is enabled only by an explicit development fallback flag.
    """

    def __init__(
        self,
        chunking: ChunkingConfig | None = None,
        *,
        tokenizer: TokenCounter | None = None,
        max_pages: int = 500,
        max_file_bytes: int = 100 * 1024 * 1024,
    ) -> None:
        self.chunking = chunking or ChunkingConfig()
        if tokenizer is None:
            from settings import RAGSettings

            settings = RAGSettings()
            settings.activate_offline_mode()
            tokenizer = (
                RegexTokenCounter()
                if settings.allow_fallback
                else LocalTokenizer(settings.embedding_model_path)
            )
        self.tokenizer = tokenizer
        if max_pages < 1 or max_file_bytes < 1:
            raise ValueError("Document limits must be positive")
        self.max_pages = max_pages
        self.max_file_bytes = max_file_bytes

    @classmethod
    def from_settings(cls, settings: RAGSettings) -> DocumentParser:
        """Build with the real local tokenizer unless development explicitly opts out."""
        settings.activate_offline_mode()
        tokenizer: TokenCounter
        if settings.mode == "development" and settings.allow_fallback:
            tokenizer = RegexTokenCounter()
        else:
            tokenizer = LocalTokenizer(settings.embedding_model_path)
        return cls(
            ChunkingConfig(
                target_tokens=settings.chunk_target_tokens,
                max_tokens=settings.chunk_max_tokens,
                overlap_tokens=settings.chunk_overlap_tokens,
            ),
            tokenizer=tokenizer,
            max_pages=settings.max_document_pages,
            max_file_bytes=settings.max_upload_file_size_mb * 1024 * 1024,
        )

    def parse(self, path: str | Path, *, backend: str = "auto") -> ParsedDocument:
        """Parse a UTF-8 text, native DOCX or text-bearing PDF; fail on unsupported input."""
        file_path = Path(path).resolve(strict=True)
        if not file_path.is_file() or not 0 < file_path.stat().st_size <= self.max_file_bytes:
            raise DocumentParseError("Document must be a nonempty file within the size limit")
        suffix = file_path.suffix.lower()
        adapters = {".txt": "text", ".md": "text", ".pdf": "pymupdf", ".docx": "docx"}
        if suffix not in adapters:
            raise DocumentParseError(f"Unsupported document extension: {suffix}")
        if backend == "auto":
            backend = adapters[suffix]
        if backend != adapters[suffix]:
            raise DocumentParseError(f"Use the local {adapters[suffix]} adapter for {suffix}")
        try:
            if backend == "text":
                elements = self._parse_lines(
                    file_path.read_text(encoding="utf-8-sig").splitlines(), str(file_path)
                )
            elif backend == "docx":
                elements = self._parse_docx(file_path)
            else:
                elements = self._parse_pymupdf(file_path)
            with file_path.open("rb") as stream:
                document_id = hashlib.file_digest(stream, "sha256").hexdigest()
            return self._finish(elements, str(file_path), document_id, backend)
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError(f"Failed to parse {file_path.name}: {exc}") from exc

    def parse_text(
        self, text: str, *, source: str = "memory", document_id: str | None = None
    ) -> ParsedDocument:
        """Parse supplied text without filesystem access."""
        if len(text.encode("utf-8")) > self.max_file_bytes:
            raise DocumentParseError("Text exceeds document size limit")
        return self._finish(
            self._parse_lines(text.splitlines(), source),
            source,
            document_id or hashlib.sha256(text.encode("utf-8")).hexdigest(),
            "text",
        )

    def _finish(
        self, elements: Iterable[DocumentElement], source: str, document_id: str, backend: str
    ) -> ParsedDocument:
        normalized = list(self._with_hierarchy(elements, source))
        if not normalized:
            raise DocumentParseError("No extractable text or tables")
        normalized = self._link_context(normalized)
        chunks = tuple(self.chunk(normalized, document_id=document_id, source=source))
        return ParsedDocument(
            document_id,
            source,
            tuple(normalized),
            chunks,
            {
                "backend": backend,
                "tokenizer": type(self.tokenizer).__name__,
                "pages_count": max((e.page or 0 for e in normalized), default=0) or None,
            },
        )

    def _parse_lines(
        self, lines: Iterable[str], source: str, page: int | None = None
    ) -> list[DocumentElement]:
        raw = list(lines)
        output: list[DocumentElement] = []
        i = 0
        while i < len(raw):
            line = raw[i].strip()
            i += 1
            if not line:
                continue
            kind, level, number, metadata = "paragraph", 0, None, {}
            md = _MARKDOWN_HEADING.match(line)
            numbered = _NUMBERED_HEADING.match(line)
            caption = _CAPTION.match(line)
            if md:
                kind, level, line = "heading", len(md[1]), md[2]
            elif _FOOTNOTE.match(line) or _MARKER.match(line):
                kind = "footnote"
            elif caption:
                kind = "caption"
                metadata = {"table_number": caption[2], "is_continuation": bool(caption[1])}
            elif numbered and len(line) < 180:
                number = numbered[1].rstrip(".")
                kind, level, line = "heading", number.count(".") + 1, numbered[2]
            elif _looks_like_table(line, raw[i] if i < len(raw) else None) or "\t" in line:
                rows = [line]
                while i < len(raw) and raw[i].strip() and ("|" in raw[i] or "\t" in raw[i]):
                    rows.append(raw[i])
                    i += 1
                kind, line = "table", _table_to_markdown(rows)
            elif re.match(r"^(?:[-•–]|[а-яa-z]\))\s", line):
                kind = "list"
            output.append(
                DocumentElement(
                    _stable_id(source, str(page), str(i), line),
                    kind,
                    line,
                    level,
                    number,
                    page=page,
                    source=source,
                    metadata=metadata,
                )
            )
        return output

    @staticmethod
    def _with_hierarchy(
        elements: Iterable[DocumentElement], source: str
    ) -> Iterator[DocumentElement]:
        stack: list[tuple[int, str]] = []
        for element in elements:
            if element.kind == "heading":
                level = max(element.level, 1)
                while stack and stack[-1][0] >= level:
                    stack.pop()
                title = f"{element.number} {element.text}" if element.number else element.text
                stack.append((level, title))
            yield replace(element, section_path=tuple(title for _, title in stack), source=source)

    @staticmethod
    def _link_context(elements: list[DocumentElement]) -> list[DocumentElement]:
        output: list[DocumentElement] = []
        caption: DocumentElement | None = None
        tables: dict[str, DocumentElement] = {}
        notes = [e for e in elements if e.kind == "footnote"]
        for element in elements:
            meta = dict(element.metadata)
            if element.kind == "heading":
                caption = None
            if element.kind == "caption":
                caption = element
            if element.kind == "table":
                columns = tuple(
                    cell.strip() for cell in element.text.splitlines()[0].strip("|").split("|")
                )
                meta.update(
                    {
                        "columns": columns,
                        "page": element.page,
                        "section_path": element.section_path,
                        "is_continuation": False,
                    }
                )
                if caption is not None:
                    meta.update(caption.metadata)
                    meta["caption"] = caption.text
                    number = str(meta["table_number"])
                    previous = tables.get(number)
                    if meta["is_continuation"]:
                        meta["continuation_of"] = previous.element_id if previous else None
                        meta["continuation_unresolved"] = previous is None
                        if previous:
                            meta["original_columns"] = previous.metadata["columns"]
                            meta["original_caption"] = previous.metadata.get("caption", "")
                    element = replace(element, metadata=meta)
                    tables[number] = element
                    caption = None
            linked: list[str] = []
            for note in notes:
                if (
                    note.element_id == element.element_id
                    or note.section_path != element.section_path
                ):
                    continue
                if note.page is not None and note.page != element.page:
                    continue
                match = _MARKER.match(note.text)
                if match:
                    marker = match[1] or match[2] or match[3]
                    references = (
                        (f"[^{marker}]", f"[{marker}]", f"{marker})")
                        if marker.isdigit()
                        else (marker,)
                    )
                    if any(reference in element.text for reference in references):
                        linked.append(note.text)
                elif note.text.lower().startswith(("примечание", "note")):
                    linked.append(note.text)
            meta["footnote_context"] = tuple(linked)
            output.append(replace(element, metadata=meta))
        return output

    def chunk(
        self, elements: Sequence[DocumentElement], *, document_id: str, source: str
    ) -> Iterator[DocumentChunk]:
        """Keep table rows atomic; count prefixes/notes and flag unavoidable overflow."""
        ordinal = 0
        pending: list[DocumentElement] = []

        def content(items: Sequence[DocumentElement]) -> str:
            path = " > ".join(items[0].section_path)
            prefix = f"Документ: {document_id}\nРаздел: {path}\n"
            notes = list(
                dict.fromkeys(
                    str(n) for item in items for n in item.metadata.get("footnote_context", ())
                )
            )
            return (
                prefix
                + "\n".join(item.text for item in items)
                + ("\n" + "\n".join(notes) if notes else "")
            )

        def emit(items: Sequence[DocumentElement]) -> DocumentChunk:
            nonlocal ordinal
            text = content(items)
            pages = [e.page for e in items if e.page is not None]
            count = self.tokenizer.count(text)
            identifier = _stable_id(document_id, str(ordinal), text)
            context = ChunkMetadata(
                document_id,
                identifier,
                items[0].section_path,
                min(pages) if pages else None,
                max(pages) if pages else None,
                tuple(e.metadata for e in items if e.kind == "table"),
                tuple(
                    dict.fromkeys(
                        str(n) for e in items for n in e.metadata.get("footnote_context", ())
                    )
                ),
                count,
            )
            ordinal += 1
            return DocumentChunk(
                identifier,
                text,
                document_id,
                source,
                context.section_path,
                context.page_start,
                context.page_end,
                tuple(dict.fromkeys(e.element_id for e in items)),
                count,
                {
                    "kinds": sorted({e.kind for e in items}),
                    "ordinal": ordinal - 1,
                    "table_context": context.table_context,
                    "footnote_context": context.footnote_context,
                    "oversized": count > self.chunking.max_tokens,
                    "tokenizer": type(self.tokenizer).__name__,
                },
            )

        for element in elements:
            if pending and (
                pending[0].section_path != element.section_path or element.kind == "table"
            ):
                yield emit(pending)
                pending = []
            if element.kind == "table":
                lines = element.text.splitlines()
                header, rows = lines[:2], lines[2:]
                batch: list[str] = []
                for row in rows:
                    candidate = replace(element, text="\n".join(header + batch + [row]))
                    if (
                        batch
                        and self.tokenizer.count(content([candidate])) > self.chunking.max_tokens
                    ):
                        yield emit([replace(element, text="\n".join(header + batch))])
                        batch = []
                    batch.append(row)
                if batch or not rows:
                    yield emit([replace(element, text="\n".join(header + batch))])
                continue
            if (
                pending
                and self.tokenizer.count(content(pending + [element])) > self.chunking.max_tokens
            ):
                yield emit(pending)
                pending = []
            if self.tokenizer.count(content([element])) <= self.chunking.max_tokens:
                pending.append(element)
                if self.tokenizer.count(content(pending)) >= self.chunking.target_tokens:
                    yield emit(pending)
                    pending = []
                continue
            # Split by character offsets, preserving punctuation and exact source substrings.
            start = 0
            while start < len(element.text):
                lo, hi = start + 1, len(element.text)
                best = start
                while lo <= hi:
                    middle = (lo + hi) // 2
                    part = replace(element, text=element.text[start:middle])
                    if self.tokenizer.count(content([part])) <= self.chunking.max_tokens:
                        best, lo = middle, middle + 1
                    else:
                        hi = middle - 1
                if best == start:
                    # Metadata or a linked note alone exceeds the budget; expose the overflow.
                    yield emit([replace(element, text=element.text[start:])])
                    break
                yield emit([replace(element, text=element.text[start:best])])
                if best == len(element.text):
                    break
                overlap_start = best
                while (
                    overlap_start > start + 1
                    and self.tokenizer.count(element.text[overlap_start - 1 : best])
                    <= self.chunking.overlap_tokens
                ):
                    overlap_start -= 1
                start = max(start + 1, overlap_start)
        if pending:
            yield emit(pending)

    def _parse_pymupdf(self, path: Path) -> list[DocumentElement]:
        import pymupdf

        output: list[DocumentElement] = []
        with pymupdf.open(path) as document:  # type: ignore[no-untyped-call]
            if document.needs_pass or document.page_count > self.max_pages:
                raise DocumentParseError("Encrypted PDF or page limit exceeded")
            margins: Counter[str] = Counter()
            for page in document:
                margin_text = {
                    str(b[4]).strip()
                    for b in page.get_text("blocks")
                    if b[1] < page.rect.height * 0.08 or b[3] > page.rect.height * 0.92
                }
                margins.update(margin_text)
            for page_index, page in enumerate(document, start=1):
                positioned: list[tuple[float, float, DocumentElement]] = []
                rects: list[Any] = []
                for index, table in enumerate(page.find_tables().tables):
                    rect = pymupdf.Rect(table.bbox)  # type: ignore[no-untyped-call]
                    rects.append(rect)
                    # Native extractor resolves geometric spans; retain raw cells for auditing.
                    rows = table.extract()
                    text = str(table.to_markdown(fill_empty=True))
                    element = DocumentElement(
                        _stable_id(str(page_index), str(index), text),
                        "table",
                        text,
                        page=page_index,
                        source=str(path),
                        metadata={
                            "bbox": (rect.x0, rect.y0, rect.x1, rect.y1),
                            "raw_cells": rows,
                            "header_names": list(table.header.names),
                        },
                    )
                    positioned.append((rect.y0, rect.x0, element))
                for index, block in enumerate(page.get_text("blocks", sort=True)):
                    if len(block) > 6 and block[6] != 0:
                        continue
                    rect = pymupdf.Rect(block[:4])  # type: ignore[no-untyped-call]
                    text = str(block[4]).strip()
                    if not text or any(
                        (rect & t).get_area() / max(rect.get_area(), 1) > 0.5 for t in rects
                    ):
                        continue
                    at_margin = (
                        rect.y0 < page.rect.height * 0.08 or rect.y1 > page.rect.height * 0.92
                    )
                    if at_margin and (
                        re.fullmatch(r"[-–—\s]*\d+[-–—\s]*", text)
                        or margins[text] >= max(2, (document.page_count + 1) // 2)
                    ):
                        continue
                    for element in self._parse_lines(
                        text.splitlines(), f"{path}#{index}", page_index
                    ):
                        positioned.append(
                            (
                                rect.y0,
                                rect.x0,
                                replace(
                                    element,
                                    metadata={
                                        **element.metadata,
                                        "bbox": (rect.x0, rect.y0, rect.x1, rect.y1),
                                    },
                                ),
                            )
                        )
                if not positioned:
                    raise DocumentParseError(
                        f"Page {page_index} has no extractable content; local OCR required"
                    )
                output.extend(item[2] for item in sorted(positioned, key=lambda item: item[:2]))
        return output

    def _parse_docx(self, path: Path) -> list[DocumentElement]:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from zipfile import ZipFile
        import xml.etree.ElementTree as ET

        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > self.max_file_bytes * 10:
                raise DocumentParseError("DOCX decompressed size limit exceeded")
            notes: dict[str, str] = {}
            if "word/footnotes.xml" in archive.namelist():
                root = ET.fromstring(archive.read("word/footnotes.xml"))
                for note in root.findall("w:footnote", namespace):
                    identifier = note.get(f"{{{namespace['w']}}}id", "")
                    if identifier.isdigit() and int(identifier) > 0:
                        notes[identifier] = " ".join(
                            n.text or "" for n in note.findall(".//w:t", namespace)
                        )
        document = Document(str(path))
        output: list[DocumentElement] = []
        for index, block in enumerate(document.iter_inner_content()):
            if isinstance(block, Paragraph):
                text = block.text
                referenced = [
                    str(node.get(f"{{{namespace['w']}}}id"))
                    for node in block._p.xpath(".//w:footnoteReference")
                ]
                for identifier in referenced:
                    if identifier not in notes:
                        raise DocumentParseError(f"Unresolved DOCX footnote {identifier}")
                    text += f" [^{identifier}]"
                style = block.style.name if block.style else ""
                heading = re.search(r"(?:Heading|Заголовок)\s*(\d+)", style or "", re.I)
                if heading and text.strip():
                    output.append(
                        DocumentElement(
                            _stable_id(str(index), text),
                            "heading",
                            text,
                            level=int(heading[1]),
                            source=str(path),
                        )
                    )
                else:
                    output.extend(self._parse_lines(text.splitlines(), f"{path}#{index}"))
                for identifier in referenced:
                    output.append(
                        DocumentElement(
                            _stable_id(str(index), identifier),
                            "footnote",
                            f"[^{identifier}]: {notes[identifier]}",
                            source=str(path),
                        )
                    )
            elif isinstance(block, Table):
                rows = [
                    [cell.text.replace("|", "&#124;").replace("\n", "<br>") for cell in row.cells]
                    for row in block.rows
                ]
                # python-docx expands gridSpan/vMerge across the layout grid.
                header_count = 0
                for row in block.rows:
                    if row._tr.xpath("./w:trPr/w:tblHeader"):
                        header_count += 1
                    else:
                        break
                header_count = max(1, header_count)
                if header_count > 1 and rows:
                    width = max(len(row) for row in rows[:header_count])
                    columns = [
                        " / ".join(
                            dict.fromkeys(
                                row[column]
                                for row in rows[:header_count]
                                if column < len(row) and row[column]
                            )
                        )
                        for column in range(width)
                    ]
                    normalized_rows = [columns] + rows[header_count:]
                else:
                    normalized_rows = rows
                text = _table_to_markdown(["\t".join(row) for row in normalized_rows])
                if text:
                    output.append(
                        DocumentElement(
                            _stable_id(str(index), text),
                            "table",
                            text,
                            source=str(path),
                            metadata={
                                "raw_cells": rows,
                                "header_rows": header_count,
                                "pagination": "unavailable",
                            },
                        )
                    )
        return output
